import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedBlock:
    text: str
    page_number: int | None = None
    sheet_name: str | None = None
    section_title: str | None = None
    clause_reference: str | None = None


class UnsupportedOCR(Exception):
    pass


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_document(path: Path) -> list[ParsedBlock]:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return [ParsedBlock(text=path.read_text(encoding="utf-8", errors="ignore"))]
    if suffix == ".csv":
        return [ParsedBlock(text=_csv_to_markdown(path))]
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".xlsx":
        return _parse_xlsx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    raise ValueError(f"Unsupported document type: {suffix}")


def _csv_to_markdown(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        rows = list(csv.reader(handle))
    return _rows_to_markdown(rows)


def _rows_to_markdown(rows: list[list[object]]) -> str:
    clean = [[str(cell or "").strip() for cell in row] for row in rows if any(str(cell or "").strip() for cell in row)]
    if not clean:
        return ""
    width = max(len(row) for row in clean)
    clean = [row + [""] * (width - len(row)) for row in clean]
    header = clean[0]
    sep = ["---"] * width
    body = clean[1:]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(sep) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _parse_docx(path: Path) -> list[ParsedBlock]:
    from docx import Document as DocxDocument

    document = DocxDocument(path)
    blocks: list[ParsedBlock] = []
    current_title = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            current_title = text
        blocks.append(ParsedBlock(text=text, section_title=current_title, clause_reference=_clause_ref(text)))
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        blocks.append(ParsedBlock(text=_rows_to_markdown(rows), section_title=current_title))
    return blocks


def _parse_xlsx(path: Path) -> list[ParsedBlock]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True)
    blocks: list[ParsedBlock] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        blocks.append(ParsedBlock(text=_rows_to_markdown(rows), sheet_name=sheet.title))
    return blocks


def _parse_pdf(path: Path) -> list[ParsedBlock]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks = []
    total = 0
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        total += len(text.strip())
        blocks.append(ParsedBlock(text=text, page_number=index))
    if total < 100:
        raise UnsupportedOCR("OCR_NOT_SUPPORTED_IN_MVP")
    return blocks


def split_into_chunks(blocks: list[ParsedBlock], target_chars: int = 3200) -> list[ParsedBlock]:
    chunks: list[ParsedBlock] = []
    for block in blocks:
        parts = _split_text(block.text, target_chars)
        for part in parts:
            chunks.append(
                ParsedBlock(
                    text=part,
                    page_number=block.page_number,
                    sheet_name=block.sheet_name,
                    section_title=block.section_title or _title_guess(part),
                    clause_reference=block.clause_reference or _clause_ref(part),
                )
            )
    return chunks


def _split_text(text: str, target_chars: int) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not text:
        return []
    clause_parts = re.split(r"(?=\n?\d+(?:\.\d+)+\s+)", text)
    paragraphs = clause_parts if len(clause_parts) > 1 else text.split("\n\n")
    chunks: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(buffer) + len(paragraph) > target_chars and buffer:
            chunks.append(buffer.strip())
            buffer = paragraph
        else:
            buffer = f"{buffer}\n\n{paragraph}".strip()
    if buffer:
        chunks.append(buffer.strip())
    return chunks


def _clause_ref(text: str) -> str | None:
    match = re.match(r"\s*(\d+(?:\.\d+)+)", text)
    return match.group(1) if match else None


def _title_guess(text: str) -> str | None:
    first = text.strip().splitlines()[0] if text.strip() else ""
    return first[:120] if len(first) < 160 else None
